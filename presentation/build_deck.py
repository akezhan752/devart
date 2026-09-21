import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY = RGBColor(0x16, 0x23, 0x3F)
NAVY_2 = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x1E, 0x93, 0xAB)
TEAL_LT = RGBColor(0x6F, 0xD6, 0xE8)
ICE = RGBColor(0xEA, 0xF1, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x1C, 0x27, 0x33)
MUTED = RGBColor(0x5B, 0x6B, 0x7B)
MUTED_LT = RGBColor(0x9F, 0xB4, 0xC2)
BORDER = RGBColor(0xD7, 0xE3, 0xEC)

HEAD = "Cambria"
BODY = "Calibri"

W = 13.333
H = 7.5
M = 0.7


def add_slide(prs, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY if dark else WHITE
    return s


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def para(tf, text, size, color, bold=False, font=BODY, space_after=0,
         align=PP_ALIGN.LEFT, first=False, italic=False, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    if line:
        p.line_spacing = line
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = font
    f.color.rgb = color
    return p


def card(slide, x, y, w, h, fill=ICE, line=None, radius=0.06):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def circle(slide, x, y, d, label, fill=TEAL, fg=WHITE, size=14):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = BODY
    r.font.color.rgb = fg
    return sh


def kicker(slide, x, y, text, color=TEAL):
    tf = textbox(slide, x, y, 8, 0.28)
    para(tf, text.upper(), 11, color, bold=True, first=True)


def title(slide, x, y, w, text, color=NAVY, size=30):
    tf = textbox(slide, x, y, w, 1.0)
    para(tf, text, size, color, bold=True, font=HEAD, first=True, line=1.05)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)

# ---------------------------------------------------------------- slide 1
s = add_slide(prs, dark=True)

card(s, W - 4.9, -1.2, 6.0, 6.0, fill=NAVY_2, radius=0.5)

kicker(s, M, 0.7, "Devart × KazService · направление 3 · AI Content & Digital Product", TEAL_LT)
title(s, M, 1.1, 8.2, "Работающий прототип, а не макет:\nсайт, регистрация и CRM-логика", WHITE, 33)

tf = textbox(s, M, 2.65, 7.6, 0.9)
para(tf, "Мероприятие получает цифровой контур целиком — от первого экрана до звонка "
         "после события. Всё открывается в браузере без установки и внешних сервисов.",
     15, MUTED_LT, first=True, line=1.25)

steps = [
    ("1", "Узнал", "LinkedIn / email"),
    ("2", "Изучил", "Лендинг"),
    ("3", "Оставил заявку", "Форма + согласие"),
    ("4", "Подтверждение", "Письмо, ID заявки"),
    ("5", "Пришёл", "Отметка на входе"),
    ("6", "Follow-up", "Звонок, материалы"),
]
cx = M
cw = 1.92
gap = 0.06
top = 4.15
for i, (n, t, sub) in enumerate(steps):
    x = cx + i * (cw + gap)
    card(s, x, top, cw, 1.55, fill=NAVY_2)
    circle(s, x + 0.16, top + 0.18, 0.36, n, fill=TEAL, size=12)
    tf = textbox(s, x + 0.16, top + 0.68, cw - 0.32, 0.32)
    para(tf, t, 12.5, WHITE, bold=True, first=True)
    tf2 = textbox(s, x + 0.16, top + 1.02, cw - 0.32, 0.42)
    para(tf2, sub, 10, MUTED_LT, first=True, line=1.1)

tf = textbox(s, M, 6.15, 11.9, 0.5)
para(tf, "Путь участника — единая система: каждый шаг оставляет запись в CRM-карточке "
         "и определяет следующее действие организатора.",
     11.5, MUTED_LT, first=True, italic=True)

notes(s, "Главный результат — работающий прототип. Открывается локально без сборки, "
         "проверен в браузере: форма, валидация, CRM-доска, мобильная вёрстка.")

# ---------------------------------------------------------------- slide 2
s = add_slide(prs)

kicker(s, M, 0.62, "Структура сайта и UX")
title(s, M, 0.95, 11.9, "Одна страница ведёт к одному действию — регистрации")

tf = textbox(s, M, 1.95, 11.9, 0.4)
para(tf, "Восемь секций закрывают возражения в том порядке, в котором они возникают у "
         "корпоративного участника.", 14, MUTED, first=True)

secs = [
    ("Первый экран", "Что, где, когда + CTA", "Решение за 5 секунд"),
    ("Для кого", "4 роли участников", "Участник узнаёт себя"),
    ("Программа", "6 блоков с таймингом", "Понятно, за что отдаёт 3 часа"),
    ("Спикеры", "Devart и KazService", "Доверие к содержанию"),
    ("Как это устроено", "Анимированная схема", "Прозрачность после заявки"),
    ("Площадка", "Формат, язык, дресс-код", "Снимает бытовые вопросы"),
    ("FAQ", "4 частых вопроса", "Убирает переписку"),
    ("Регистрация", "Форма + согласие", "Целевое действие"),
]
cw = 2.86
ch = 1.42
gapx = 0.19
gapy = 0.19
x0 = M
y0 = 2.55
for i, (t, what, why) in enumerate(secs):
    col = i % 4
    row = i // 4
    x = x0 + col * (cw + gapx)
    y = y0 + row * (ch + gapy)
    card(s, x, y, cw, ch, fill=ICE)
    tf = textbox(s, x + 0.22, y + 0.19, cw - 0.44, 0.3)
    para(tf, t, 13, NAVY, bold=True, first=True)
    tf = textbox(s, x + 0.22, y + 0.53, cw - 0.44, 0.3)
    para(tf, what, 10.5, MUTED, first=True)
    tf = textbox(s, x + 0.22, y + 0.92, cw - 0.44, 0.35)
    para(tf, why, 10.5, TEAL, bold=True, first=True)

card(s, M, 5.95, 11.9, 0.95, fill=NAVY)
tf = textbox(s, M + 0.3, 6.15, 11.3, 0.6)
para(tf, "Что это значит для бизнеса", 11, TEAL_LT, bold=True, first=True)
para(tf, "Структура держит внимание до формы: нижняя строка каждой карточки — не описание, "
         "а снятое возражение. Мобильная вёрстка проверена — горизонтальной прокрутки нет.",
     12, WHITE, space_after=0)

notes(s, "UX-принцип: одна страница — одно целевое действие. Секции отвечают на возражения "
         "в порядке их возникновения.")

# ---------------------------------------------------------------- slide 3
s = add_slide(prs)

kicker(s, M, 0.62, "Регистрация, CRM и схема данных")
title(s, M, 0.95, 11.9, "Заявка сразу становится карточкой со статусом и следующим шагом")

left_w = 5.5
tf = textbox(s, M, 1.95, left_w, 0.3)
para(tf, "Форма", 13, NAVY, bold=True, first=True)

fields = [
    "Имя и фамилия — обязательно",
    "Компания — обязательно",
    "Роль — список из 5 значений",
    "Email — обязательно, формат проверяется",
    "Телефон — необязательно, формат проверяется",
    "Согласие на обработку ПД — обязательно",
]
y = 2.35
for f in fields:
    circle(s, M, y + 0.03, 0.16, "", fill=TEAL)
    tf = textbox(s, M + 0.3, y, left_w - 0.3, 0.28)
    para(tf, f, 12, TEXT, first=True)
    y += 0.34

card(s, M, 4.55, left_w, 1.15, fill=ICE)
tf = textbox(s, M + 0.25, 4.75, left_w - 0.5, 0.8)
para(tf, "Валидация — на каждом поле", 11.5, NAVY, bold=True, first=True)
para(tf, "Ошибка появляется при потере фокуса и исчезает по мере исправления. "
         "Без согласия отправка не проходит.", 11, MUTED, space_after=0, line=1.15)

rx = M + left_w + 0.55
tf = textbox(s, rx, 1.95, 11.9 - left_w - 0.55, 0.3)
para(tf, "Статусы участника", 13, NAVY, bold=True, first=True)

statuses = [
    ("Заявка", "форма отправлена и прошла проверку"),
    ("Подтверждено", "письмо с деталями участия"),
    ("Напоминание", "за сутки до мероприятия"),
    ("Посещение", "отметка на входе"),
    ("Follow-up", "звонок и материалы после"),
]
y = 2.35
for i, (st, desc) in enumerate(statuses):
    circle(s, rx, y - 0.02, 0.3, str(i + 1), fill=TEAL, size=11)
    tf = textbox(s, rx + 0.45, y, 2.1, 0.28)
    para(tf, st, 12, NAVY, bold=True, first=True)
    tf = textbox(s, rx + 2.5, y, 3.6, 0.3)
    para(tf, desc, 11, MUTED, first=True)
    y += 0.44

card(s, rx, 4.55, 11.9 - left_w - 0.55, 1.15, fill=ICE)
tf = textbox(s, rx + 0.25, 4.75, 11.9 - left_w - 1.05, 0.8)
para(tf, "Схема записи", 11.5, NAVY, bold=True, first=True)
para(tf, "id · name · company · role · email · phone · status · createdAt", 11, MUTED,
     space_after=0, font="Courier New")

card(s, M, 5.95, 11.9, 0.95, fill=NAVY)
tf = textbox(s, M + 0.3, 6.15, 11.3, 0.6)
para(tf, "Что это значит для бизнеса", 11, TEAL_LT, bold=True, first=True)
para(tf, "Руководитель видит воронку целиком: сколько заявок, сколько подтвердились, "
         "сколько дошли. В прототипе переходы кнопкой — в проде это триггеры почты, QR и задача менеджеру.",
     12, WHITE, space_after=0)

notes(s, "Логика: форма → проверка → CRM → подтверждение → напоминание → посещение → follow-up. "
         "Хранилище — localStorage, только тестовые данные.")

# ---------------------------------------------------------------- slide 4
s = add_slide(prs)

kicker(s, M, 0.62, "AI-инструменты и ручная доработка")
title(s, M, 0.95, 11.9, "AI дал черновик и скорость — правки решали качество")

tf = textbox(s, M, 1.95, 11.9, 0.4)
para(tf, "Всё, что AI сгенерировал, проверялось в браузере. Ниже — что именно пришлось исправить руками.",
     14, MUTED, first=True)

rows = [
    ("Вёрстка и разметка", "Каркас страниц, сетки, адаптив",
     "Анимация SVG через атрибут r не работает в Safari — заменил на transform"),
    ("Тексты на русском", "Черновики секций, FAQ, формулировки",
     "Счётчик выдавал «1 заявок» — добавил правило склонения (1/2–4/5+)"),
    ("Логика формы и CRM", "Валидация, статусы, хранение",
     "Демо-данные покрывали 4 статуса из 5 — добавил пятую запись"),
    ("Структура подачи", "План слайдов и документов",
     "Переписал заголовки слайдов: тезис вместо названия раздела"),
]
y = 2.6
rh = 0.92
for i, (area, ai, manual) in enumerate(rows):
    fill = ICE if i % 2 == 0 else WHITE
    card(s, M, y, 11.9, rh, fill=fill)
    tf = textbox(s, M + 0.3, y + 0.16, 2.7, 0.6)
    para(tf, area, 12.5, NAVY, bold=True, first=True)
    tf = textbox(s, M + 3.15, y + 0.16, 3.5, 0.62)
    para(tf, ai, 11.5, MUTED, first=True, line=1.15)
    circle(s, M + 6.85, y + 0.3, 0.26, "!", fill=TEAL, size=11)
    tf = textbox(s, M + 7.25, y + 0.16, 4.35, 0.66)
    para(tf, manual, 11.5, TEXT, first=True, line=1.15)
    y += rh + 0.1

tf = textbox(s, M + 3.15, 2.3, 3.5, 0.26)
para(tf, "ЧТО СДЕЛАЛ AI", 9.5, MUTED, bold=True, first=True)
tf = textbox(s, M + 7.25, 2.3, 4.35, 0.26)
para(tf, "ЧТО ИСПРАВИЛ ВРУЧНУЮ", 9.5, TEAL, bold=True, first=True)

tf = textbox(s, M, 6.5, 11.9, 0.4)
para(tf, "Инструменты: Claude (код, тексты, структура) · проверка — ручное тестирование в браузере "
         "на десктопе и мобильном разрешении.", 11.5, MUTED, first=True, italic=True)

notes(s, "Ключевая мысль: необработанная генерация не считается результатом. Каждый пункт справа — "
         "дефект, найденный при ручной проверке.")

# ---------------------------------------------------------------- slide 5
s = add_slide(prs, dark=True)

kicker(s, M, 0.62, "Ограничения, риски и следующий backlog", TEAL_LT)
title(s, M, 0.95, 11.9, "Прототип честно показывает логику — но это ещё не продакшн",
      WHITE, 29)

col_w = 3.75
gapc = 0.32

blocks = [
    ("Ограничения", NAVY_2, [
        "Данные — в localStorage браузера, не в базе",
        "Письма и напоминания не отправляются",
        "Нет сервера: защита и доступы не реализованы",
        "Спикеры и адрес площадки — плейсхолдеры",
    ]),
    ("Главные риски", NAVY_2, [
        "Демо примут за готовую систему",
        "Данные исчезнут при очистке браузера",
        "Тексты не проверены юристом Devart",
        "Реальная CRM может не принять эту схему",
    ]),
    ("Следующий backlog", NAVY_2, [
        "Форма → Make/n8n → таблица или CRM",
        "Автописьмо-подтверждение и напоминание",
        "QR-check-in на входе",
        "Казахская версия страницы",
    ]),
]
for i, (head, fill, items) in enumerate(blocks):
    x = M + i * (col_w + gapc)
    card(s, x, 2.0, col_w, 3.5, fill=fill)
    tf = textbox(s, x + 0.28, 2.25, col_w - 0.56, 0.32)
    para(tf, head, 14, TEAL_LT, bold=True, first=True)
    yy = 2.78
    for it in items:
        circle(s, x + 0.28, yy + 0.06, 0.13, "", fill=TEAL)
        tf = textbox(s, x + 0.55, yy, col_w - 0.83, 0.62)
        para(tf, it, 11.5, WHITE, first=True, line=1.15)
        yy += 0.65

card(s, M, 5.75, 11.9, 1.05, fill=TEAL)
tf = textbox(s, M + 0.35, 5.98, 11.2, 0.7)
para(tf, "Первый шаг", 11, NAVY, bold=True, first=True)
para(tf, "Подключить форму к таблице через Make и включить автописьмо — после этого прототип "
         "можно ставить на реальный поток регистраций без переписывания фронтенда.",
     12.5, WHITE, space_after=0)

notes(s, "Главный риск — принять демо за систему. Первый шаг делает прототип пригодным для "
         "реального потока без переделки фронтенда.")

out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "Devart_KazService_Digital_Product.pptx")
prs.save(out)
print("saved:", out)
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
