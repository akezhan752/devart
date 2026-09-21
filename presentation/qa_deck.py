import sys
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400
CHAR_W = {"Calibri": 0.47, "Cambria": 0.50, "Courier New": 0.60}


def est_lines(text, pt_size, font, width_in):
    if not text.strip():
        return 0
    w = CHAR_W.get(font, 0.48) * pt_size / 72.0
    per_line = max(1, int(width_in / w))
    lines = 0
    for chunk in text.split("\n"):
        lines += max(1, -(-len(chunk) // per_line))
    return lines


def main(path):
    prs = Presentation(path)
    SW = prs.slide_width / EMU_IN
    SH = prs.slide_height / EMU_IN
    print("slide size: %.2f x %.2f in, slides: %d\n" % (SW, SH, len(prs.slides._sldIdLst)))

    problems = []
    for idx, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.left is None:
                continue
            L = sh.left / EMU_IN
            T = sh.top / EMU_IN
            W = sh.width / EMU_IN
            H = sh.height / EMU_IN
            if L < -0.01 or T < -0.01 or L + W > SW + 0.01 or T + H > SH + 0.01:
                # intentional bleed shapes are allowed; flag only text-bearing ones
                has_text = sh.has_text_frame and sh.text_frame.text.strip()
                tag = "TEXT" if has_text else "decor"
                if has_text:
                    problems.append(
                        "s%d %s OUT OF BOUNDS: L=%.2f T=%.2f R=%.2f B=%.2f  %r"
                        % (idx, tag, L, T, L + W, T + H, sh.text_frame.text[:40]))
            if L < 0.5 and L >= 0 and sh.has_text_frame and sh.text_frame.text.strip():
                if L < 0.49:
                    problems.append("s%d MARGIN: left=%.2f  %r"
                                    % (idx, L, sh.text_frame.text[:40]))

            if not sh.has_text_frame:
                continue
            total_h = 0.0
            for p in sh.text_frame.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if not txt:
                    continue
                size = None
                font = "Calibri"
                for r in p.runs:
                    if r.font.size:
                        size = r.font.size.pt
                    if r.font.name:
                        font = r.font.name
                    break
                size = size or 18
                ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
                n = est_lines(txt, size, font, W)
                total_h += n * size * ls / 72.0
                sa = p.space_after.pt / 72.0 if p.space_after else 0
                total_h += sa
            if total_h > H + 0.06:
                problems.append(
                    "s%d OVERFLOW: needs %.2f\" has %.2f\"  %r"
                    % (idx, total_h, H, sh.text_frame.text[:50]))

    if problems:
        print("ISSUES (%d):" % len(problems))
        for p in problems:
            print("  " + p)
    else:
        print("no bounds/overflow issues detected")
    return len(problems)


if __name__ == "__main__":
    sys.exit(1 if main(sys.argv[1]) else 0)
